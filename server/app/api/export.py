import csv
import io
from datetime import date, datetime

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select

from app.core.dependencies import CurrentUser, DbSession
from app.models.daily_record import DailyRecord
from app.models.gb_record import GbRecord
from app.models.guestbook import Guestbook
from app.models.operation_log import OperationLog

router = APIRouter(prefix="/export", tags=["export"])

EVENT_TYPE_MAP = {
    "wedding": "婚礼",
    "engagement": "订婚",
    "birthday": "生日",
    "fullmoon": "满月酒",
    "birthday_banquet": "生日/寿宴",
    "housewarming": "乔迁",
    "graduation": "升学",
    "funeral": "丧葬",
    "holiday": "节日红包",
    "other": "其他",
}

PAYMENT_METHOD_MAP = {
    "cash": "现金",
    "wechat": "微信",
    "other": "其他",
}

RECORD_TYPE_MAP = {
    "in": "收",
    "out": "送",
}


def _display_event_type(val: str | None) -> str:
    if not val:
        return ""
    return EVENT_TYPE_MAP.get(val, val)


def _display_payment_method(val: str | None) -> str:
    if not val:
        return ""
    return PAYMENT_METHOD_MAP.get(val, val)


def _display_record_type(val: str | None) -> str:
    if not val:
        return ""
    return RECORD_TYPE_MAP.get(val, val)


def _format_amount(cents: int) -> str:
    return f"{cents / 100:.2f}"


def _content_disposition(filename: str) -> str:
    from urllib.parse import quote
    ascii_name = filename.encode("ascii", "replace").decode("ascii")
    encoded = quote(filename)
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{encoded}"


def _build_excel_response(content: bytes, filename: str) -> StreamingResponse:
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(filename)},
    )


def _build_csv_response(content: str, filename: str) -> StreamingResponse:
    return StreamingResponse(
        io.BytesIO(content.encode("utf-8-sig")),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": _content_disposition(filename)},
    )


def _generate_daily_records_excel(records: list[DailyRecord]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "日常记录"

    headers = ["姓名", "电话", "类型", "金额(元)", "事由", "支付方式", "日期", "备注"]
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, rec in enumerate(records, 2):
        ws.cell(row=row_idx, column=1, value=rec.person_name)
        ws.cell(row=row_idx, column=2, value=rec.phone or "")
        ws.cell(row=row_idx, column=3, value=_display_record_type(rec.type))
        ws.cell(row=row_idx, column=4, value=float(rec.amount) / 100)
        ws.cell(row=row_idx, column=5, value=_display_event_type(rec.event_type))
        ws.cell(row=row_idx, column=6, value=_display_payment_method(rec.payment_method))
        ws.cell(row=row_idx, column=7, value=rec.event_date.isoformat())
        ws.cell(row=row_idx, column=8, value=rec.note or "")

    col_widths = [12, 15, 6, 12, 10, 10, 12, 30]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generate_daily_records_csv(records: list[DailyRecord]) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["姓名", "电话", "类型", "金额(元)", "事由", "支付方式", "日期", "备注"])
    for rec in records:
        writer.writerow([
            rec.person_name,
            rec.phone or "",
            _display_record_type(rec.type),
            _format_amount(rec.amount),
            _display_event_type(rec.event_type),
            _display_payment_method(rec.payment_method),
            rec.event_date.isoformat(),
            rec.note or "",
        ])
    return output.getvalue()


def _generate_guestbook_excel(
    guestbook: Guestbook, records: list[GbRecord], logs: list[OperationLog]
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()

    ws_records = wb.active
    ws_records.title = "收礼记录"

    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")

    rec_headers = ["宾客姓名", "电话", "金额(元)", "支付方式", "收银人", "记账时间", "对账状态", "备注"]
    for col, header in enumerate(rec_headers, 1):
        cell = ws_records.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, rec in enumerate(records, 2):
        ws_records.cell(row=row_idx, column=1, value=rec.guest_name)
        ws_records.cell(row=row_idx, column=2, value=rec.phone or "")
        ws_records.cell(row=row_idx, column=3, value=float(rec.amount) / 100)
        ws_records.cell(row=row_idx, column=4, value=_display_payment_method(rec.payment_method))
        ws_records.cell(row=row_idx, column=5, value=rec.agent_name or "")
        ws_records.cell(row=row_idx, column=6, value=rec.record_time.strftime("%Y-%m-%d %H:%M"))
        ws_records.cell(row=row_idx, column=7, value="已对账" if rec.check_status == "checked" else "未对账")
        ws_records.cell(row=row_idx, column=8, value=rec.note or "")

    rec_widths = [12, 15, 12, 10, 10, 18, 10, 30]
    for i, width in enumerate(rec_widths, 1):
        ws_records.column_dimensions[ws_records.cell(row=1, column=i).column_letter].width = width

    if logs:
        ws_logs = wb.create_sheet("操作日志")
        log_headers = ["操作类型", "操作人", "终端", "操作时间", "备注"]
        for col, header in enumerate(log_headers, 1):
            cell = ws_logs.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_idx, log in enumerate(logs, 2):
            ws_logs.cell(row=row_idx, column=1, value=log.action_type)
            ws_logs.cell(row=row_idx, column=2, value=log.operator_name or "")
            ws_logs.cell(row=row_idx, column=3, value=log.client_id or "")
            ws_logs.cell(row=row_idx, column=4, value=log.created_at.strftime("%Y-%m-%d %H:%M"))
            ws_logs.cell(row=row_idx, column=5, value=log.note or "")

        log_widths = [10, 10, 12, 18, 30]
        for i, width in enumerate(log_widths, 1):
            ws_logs.column_dimensions[ws_logs.cell(row=1, column=i).column_letter].width = width

    ws_summary = wb.create_sheet("汇总")
    total = sum(r.amount for r in records)
    summary_data = [
        ("礼簿名称", guestbook.name),
        ("事由", _display_event_type(guestbook.event_type)),
        ("宴席日期", guestbook.event_date.isoformat()),
        ("状态", guestbook.status),
        ("记录总数", len(records)),
        ("总金额(元)", _format_amount(total)),
    ]
    for row_idx, (label, value) in enumerate(summary_data, 1):
        cell_label = ws_summary.cell(row=row_idx, column=1, value=label)
        cell_label.font = Font(bold=True)
        ws_summary.cell(row=row_idx, column=2, value=str(value))
    ws_summary.column_dimensions["A"].width = 12
    ws_summary.column_dimensions["B"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/daily-records")
async def export_daily_records(
    current_user: CurrentUser,
    db: DbSession,
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    type: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
    include_private: bool = Query(True),
):
    stmt = (
        select(DailyRecord)
        .where(DailyRecord.owner_id == current_user.id, DailyRecord.is_deleted == False)
    )
    if not include_private:
        stmt = stmt.where(DailyRecord.is_private == False)
    if type:
        stmt = stmt.where(DailyRecord.type == type)
    if start_date:
        stmt = stmt.where(DailyRecord.event_date >= start_date)
    if end_date:
        stmt = stmt.where(DailyRecord.event_date <= end_date)
    stmt = stmt.order_by(DailyRecord.event_date.desc())

    result = await db.execute(stmt)
    records = result.scalars().all()

    if format == "csv":
        content = _generate_daily_records_csv(list(records))
        return _build_csv_response(content, f"日常记录_{date.today().isoformat()}.csv")
    else:
        content = _generate_daily_records_excel(list(records))
        return _build_excel_response(content, f"日常记录_{date.today().isoformat()}.xlsx")


@router.get("/guestbooks/{guestbook_id}")
async def export_guestbook(
    guestbook_id: int,
    current_user: CurrentUser,
    db: DbSession,
):
    result = await db.execute(
        select(Guestbook)
        .where(Guestbook.id == guestbook_id, Guestbook.owner_id == current_user.id,
               Guestbook.is_deleted == False)
    )
    guestbook = result.scalar_one_or_none()
    if not guestbook:
        raise HTTPException(status_code=404, detail="礼簿不存在")

    records_result = await db.execute(
        select(GbRecord)
        .where(GbRecord.guestbook_id == guestbook_id, GbRecord.owner_id == current_user.id,
               GbRecord.is_deleted == False)
        .order_by(GbRecord.record_time.desc())
    )
    records = records_result.scalars().all()

    logs_result = await db.execute(
        select(OperationLog)
        .where(OperationLog.owner_id == current_user.id,
               OperationLog.guestbook_id == guestbook_id)
        .order_by(OperationLog.created_at.desc())
    )
    logs = logs_result.scalars().all()

    content = _generate_guestbook_excel(guestbook, list(records), list(logs))
    filename = f"礼簿_{guestbook.name}_{date.today().isoformat()}.xlsx"
    return _build_excel_response(content, filename)
