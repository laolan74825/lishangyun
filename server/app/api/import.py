import csv
import io
import re
from datetime import date, datetime
from typing import Any

from fastapi import APIRouter, HTTPException, UploadFile, File, Path, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select

from app.core.dependencies import CurrentUser, DbSession
from app.core.helpers import write_audit_log
from app.models.daily_record import DailyRecord
from app.models.gb_record import GbRecord
from app.models.guestbook import Guestbook
from app.schemas.export_import import (
    ImportCommitRequest,
    ImportCommitResponse,
    ImportPreviewResponse,
    ImportPreviewRow,
    ImportType,
)
from app.utils.idempotency import new_uuid

router = APIRouter(prefix="/import", tags=["import"])

VALID_DAILY_EVENT_TYPES = {"婚礼", "生日", "满月酒", "乔迁", "升学", "丧葬", "节日红包", "其他"}
VALID_PAYMENT_METHODS = {"现金", "微信", "其他"}
VALID_RECORD_TYPES = {"收", "送"}

EVENT_TYPE_TO_EN = {
    "婚礼": "wedding",
    "生日": "birthday",
    "满月酒": "fullmoon",
    "乔迁": "housewarming",
    "升学": "graduation",
    "丧葬": "funeral",
    "节日红包": "holiday",
    "其他": "other",
}

PAYMENT_METHOD_TO_EN = {
    "现金": "cash",
    "微信": "wechat",
    "其他": "other",
}

RECORD_TYPE_TO_EN = {
    "收": "in",
    "送": "out",
}

DAILY_RECORD_HEADERS = ["姓名", "电话", "类型", "金额", "事由", "支付方式", "日期", "备注"]
GB_RECORD_HEADERS = ["宾客姓名", "电话", "金额", "支付方式", "收银人", "记账时间", "备注"]

MAX_ROWS = 10000
MAX_FILE_SIZE = 5 * 1024 * 1024

_preview_cache: dict[str, dict] = {}


def _build_excel_response(content: bytes, filename: str) -> StreamingResponse:
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _generate_daily_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "日常记录模板"

    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")

    for col, header in enumerate(DAILY_RECORD_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    example = ["张三", "13800138000", "收", "500.00", "婚礼", "微信", "2026-10-01", "朋友婚礼"]
    for col, val in enumerate(example, 1):
        ws.cell(row=2, column=col, value=val)

    ws_note = wb.create_sheet("字段说明")
    notes = [
        ("字段", "必填", "说明"),
        ("姓名", "是", "最多50字"),
        ("电话", "否", "最多20字"),
        ("类型", "是", "收 或 送"),
        ("金额", "是", "正数，最多2位小数，范围1-100000元"),
        ("事由", "是", "婚礼/生日/满月酒/乔迁/升学/丧葬/节日红包/其他"),
        ("支付方式", "否", "现金/微信/其他"),
        ("日期", "是", "YYYY-MM-DD 格式，不晚于今天"),
        ("备注", "否", "最多500字"),
    ]
    for row_idx, (field, required, desc) in enumerate(notes, 1):
        ws_note.cell(row=row_idx, column=1, value=field)
        ws_note.cell(row=row_idx, column=2, value=required)
        ws_note.cell(row=row_idx, column=3, value=desc)
    ws_note.column_dimensions["A"].width = 10
    ws_note.column_dimensions["B"].width = 6
    ws_note.column_dimensions["C"].width = 45

    widths = [12, 15, 6, 12, 10, 10, 12, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _generate_gb_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "礼簿记录模板"

    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")

    for col, header in enumerate(GB_RECORD_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    example = ["李四", "13900139000", "600.00", "现金", "王五", "2026-10-01 12:30", "老同学"]
    for col, val in enumerate(example, 1):
        ws.cell(row=2, column=col, value=val)

    ws_note = wb.create_sheet("字段说明")
    notes = [
        ("字段", "必填", "说明"),
        ("宾客姓名", "是", "最多50字"),
        ("电话", "否", "最多20字"),
        ("金额", "是", "正数，最多2位小数，范围1-100000元"),
        ("支付方式", "否", "现金/微信/其他"),
        ("收银人", "否", "最多50字"),
        ("记账时间", "否", "YYYY-MM-DD HH:MM 格式，默认当前时间"),
        ("备注", "否", "最多500字"),
    ]
    for row_idx, (field, required, desc) in enumerate(notes, 1):
        ws_note.cell(row=row_idx, column=1, value=field)
        ws_note.cell(row=row_idx, column=2, value=required)
        ws_note.cell(row=row_idx, column=3, value=desc)
    ws_note.column_dimensions["A"].width = 10
    ws_note.column_dimensions["B"].width = 6
    ws_note.column_dimensions["C"].width = 45

    widths = [12, 15, 12, 10, 10, 18, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/templates/{template_type}")
async def download_template(
    current_user: CurrentUser,
    template_type: str = Path(..., pattern="^(daily_record|gb_record)$"),
):
    if template_type == "daily_record":
        content = _generate_daily_template()
        return _build_excel_response(content, "日常记录导入模板.xlsx")
    else:
        content = _generate_gb_template()
        return _build_excel_response(content, "礼簿记录导入模板.xlsx")


async def _read_uploaded_file(file: UploadFile) -> tuple[list[dict], str]:
    raw = await file.read()
    if len(raw) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="导入文件过大，请分批导入（最大5MB）")

    filename = file.filename or "upload.xlsx"

    if filename.endswith(".csv"):
        text = None
        for encoding in ("utf-8-sig", "utf-8", "gbk", "gb2312"):
            try:
                text = raw.decode(encoding)
                break
            except (UnicodeDecodeError, LookupError):
                continue
        if text is None:
            raise HTTPException(status_code=400, detail="文件编码无法识别，请转换为 UTF-8 后重试")

        reader = csv.DictReader(io.StringIO(text))
        rows = [dict(row) for row in reader]
        return rows, filename

    elif filename.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="无法读取 Excel 文件，请检查文件格式")

        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            headers = [str(h).strip() if h else "" for h in next(rows_iter)]
        except StopIteration:
            raise HTTPException(status_code=400, detail="文件为空")

        rows = []
        for row_values in rows_iter:
            if all(v is None or str(v).strip() == "" for v in row_values):
                continue
            row_dict = {}
            for col_idx, header in enumerate(headers):
                if header and col_idx < len(row_values):
                    val = row_values[col_idx]
                    row_dict[header] = val if val is not None else ""
            if row_dict:
                rows.append(row_dict)
        wb.close()
        return rows, filename
    else:
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 和 .csv 格式")


def _parse_amount(val: Any) -> int | None:
    if val is None or str(val).strip() == "":
        return None
    s = str(val).strip()
    if not re.match(r"^\d+(\.\d{1,2})?$", s):
        return None
    cents = round(float(s) * 100)
    if cents < 100 or cents > 10000000:
        return None
    return cents


def _parse_date(val: Any) -> date | None:
    if val is None or str(val).strip() == "":
        return None
    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            d = datetime.strptime(s, fmt).date()
            if d > date.today():
                return None
            return d
        except ValueError:
            continue
    return None


def _parse_datetime(val: Any) -> datetime | None:
    if val is None or str(val).strip() == "":
        return None
    s = str(val).strip()
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def _validate_daily_row(row: dict, row_number: int) -> tuple[dict, list[str]]:
    data: dict[str, Any] = {}
    errors: list[str] = []

    name = str(row.get("姓名", "")).strip()
    if not name:
        errors.append(f"第 {row_number} 行缺少必填字段：姓名")
    elif len(name) > 50:
        errors.append(f"第 {row_number} 行姓名过长（最多50字）")
    else:
        data["person_name"] = name

    phone = str(row.get("电话", "")).strip()
    if phone:
        if len(phone) > 20:
            errors.append(f"第 {row_number} 行电话过长（最多20字）")
        else:
            data["phone"] = phone

    record_type = str(row.get("类型", "")).strip()
    if not record_type:
        errors.append(f"第 {row_number} 行缺少必填字段：类型")
    elif record_type not in VALID_RECORD_TYPES:
        errors.append(f"第 {row_number} 行类型无效，应为「收」或「送」")
    else:
        data["type"] = RECORD_TYPE_TO_EN[record_type]

    amount_cents = _parse_amount(row.get("金额"))
    if amount_cents is None:
        errors.append(f"第 {row_number} 行金额格式错误，应为 1-100000 之间的数字")
    else:
        data["amount"] = amount_cents

    event_type = str(row.get("事由", "")).strip()
    if not event_type:
        errors.append(f"第 {row_number} 行缺少必填字段：事由")
    elif event_type not in VALID_DAILY_EVENT_TYPES:
        errors.append(f"第 {row_number} 行事由类型无效，请从下拉列表选择")
    else:
        data["event_type"] = EVENT_TYPE_TO_EN[event_type]

    payment = str(row.get("支付方式", "")).strip()
    if payment:
        if payment not in VALID_PAYMENT_METHODS:
            errors.append(f"第 {row_number} 行支付方式无效")
        else:
            data["payment_method"] = PAYMENT_METHOD_TO_EN[payment]

    event_date = _parse_date(row.get("日期"))
    if event_date is None:
        errors.append(f"第 {row_number} 行日期格式错误或超出合理范围")
    else:
        data["event_date"] = event_date

    note = str(row.get("备注", "")).strip()
    if note:
        data["note"] = note[:500]

    return data, errors


def _validate_gb_row(row: dict, row_number: int) -> tuple[dict, list[str]]:
    data: dict[str, Any] = {}
    errors: list[str] = []

    name = str(row.get("宾客姓名", "")).strip()
    if not name:
        errors.append(f"第 {row_number} 行缺少必填字段：宾客姓名")
    elif len(name) > 50:
        errors.append(f"第 {row_number} 行宾客姓名过长（最多50字）")
    else:
        data["guest_name"] = name

    phone = str(row.get("电话", "")).strip()
    if phone:
        if len(phone) > 20:
            errors.append(f"第 {row_number} 行电话过长（最多20字）")
        else:
            data["phone"] = phone

    amount_cents = _parse_amount(row.get("金额"))
    if amount_cents is None:
        errors.append(f"第 {row_number} 行金额格式错误，应为 1-100000 之间的数字")
    else:
        data["amount"] = amount_cents

    payment = str(row.get("支付方式", "")).strip()
    if payment:
        if payment not in VALID_PAYMENT_METHODS:
            errors.append(f"第 {row_number} 行支付方式无效")
        else:
            data["payment_method"] = PAYMENT_METHOD_TO_EN[payment]
    else:
        data["payment_method"] = "cash"

    agent = str(row.get("收银人", "")).strip()
    if agent:
        if len(agent) > 50:
            errors.append(f"第 {row_number} 行收银人过长（最多50字）")
        else:
            data["agent_name"] = agent

    record_time = _parse_datetime(row.get("记账时间"))
    if record_time:
        data["record_time"] = record_time

    note = str(row.get("备注", "")).strip()
    if note:
        data["note"] = note[:500]

    return data, errors


async def _check_daily_duplicates(
    db: DbSession, owner_id: int, data: dict
) -> tuple[bool, str | None]:
    if "person_name" not in data or "amount" not in data or "event_date" not in data:
        return False, None
    result = await db.execute(
        select(DailyRecord).where(
            DailyRecord.owner_id == owner_id,
            DailyRecord.is_deleted == False,
            DailyRecord.person_name == data["person_name"],
            DailyRecord.amount == data["amount"],
            DailyRecord.event_date == data["event_date"],
        ).limit(1)
    )
    existing = result.scalar_one_or_none()
    if existing:
        summary = f"{existing.person_name} {existing.amount / 100:.0f}元 {existing.event_date.isoformat()}"
        return True, summary
    return False, None


async def _check_gb_duplicates(
    db: DbSession, guestbook_id: int, owner_id: int, data: dict
) -> tuple[bool, str | None]:
    if "guest_name" not in data or "amount" not in data:
        return False, None
    result = await db.execute(
        select(GbRecord).where(
            GbRecord.guestbook_id == guestbook_id,
            GbRecord.owner_id == owner_id,
            GbRecord.is_deleted == False,
            GbRecord.guest_name == data["guest_name"],
            GbRecord.amount == data["amount"],
        ).limit(1)
    )
    existing = result.scalar_one_or_none()
    if existing:
        summary = f"{existing.guest_name} {existing.amount / 100:.0f}元"
        return True, summary
    return False, None


@router.post("/preview", response_model=ImportPreviewResponse)
async def preview_import(
    current_user: CurrentUser,
    db: DbSession,
    file: UploadFile = File(...),
    import_type: ImportType = Query(...),
    guestbook_id: int | None = None,
):
    if import_type == "gb_record" and not guestbook_id:
        raise HTTPException(status_code=400, detail="礼簿记录导入需指定 guestbook_id")

    if import_type == "gb_record":
        result = await db.execute(
            select(Guestbook).where(
                Guestbook.id == guestbook_id,
                Guestbook.owner_id == current_user.id,
                Guestbook.is_deleted == False,
            )
        )
        if not result.scalar_one_or_none():
            raise HTTPException(status_code=404, detail="礼簿不存在")

    rows, filename = await _read_uploaded_file(file)

    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"导入行数过多（最多{MAX_ROWS}行），请分批导入")

    preview_rows: list[ImportPreviewRow] = []
    cached_rows: list[dict] = []
    valid_count = 0
    invalid_count = 0
    duplicate_count = 0

    for idx, row in enumerate(rows):
        row_number = idx + 2

        if import_type == "daily_record":
            data, errors = _validate_daily_row(row, row_number)
        else:
            data, errors = _validate_gb_row(row, row_number)

        is_duplicate = False
        duplicate_summary = None
        if not errors:
            if import_type == "daily_record":
                is_duplicate, duplicate_summary = await _check_daily_duplicates(
                    db, current_user.id, data
                )
            else:
                is_duplicate, duplicate_summary = await _check_gb_duplicates(
                    db, guestbook_id, current_user.id, data
                )

        is_valid = len(errors) == 0
        if is_valid and not is_duplicate:
            valid_count += 1
        elif is_duplicate:
            duplicate_count += 1
        else:
            invalid_count += 1

        row_data = {k: str(v) if not isinstance(v, (date, datetime)) else v.isoformat() for k, v in row.items()}
        preview_rows.append(ImportPreviewRow(
            row_number=row_number,
            data=row_data,
            is_valid=is_valid,
            errors=errors,
            is_duplicate=is_duplicate,
            duplicate_summary=duplicate_summary,
        ))

        cached_entry: dict[str, Any] = {"is_valid": is_valid, "is_duplicate": is_duplicate}
        if is_valid and not is_duplicate:
            cached_entry["_parsed_data"] = data
        cached_rows.append(cached_entry)

    _preview_cache[filename] = {"rows": cached_rows}

    return ImportPreviewResponse(
        import_type=import_type,
        total_rows=len(rows),
        valid_rows=valid_count,
        invalid_rows=invalid_count,
        duplicate_rows=duplicate_count,
        rows=preview_rows,
        file_name=filename,
    )


@router.post("/commit", response_model=ImportCommitResponse)
async def commit_import(
    body: ImportCommitRequest,
    current_user: CurrentUser,
    db: DbSession,
):
    if body.import_type == "gb_record" and not body.guestbook_id:
        raise HTTPException(status_code=400, detail="礼簿记录导入需指定 guestbook_id")

    if body.import_type == "gb_record":
        result = await db.execute(
            select(Guestbook).where(
                Guestbook.id == body.guestbook_id,
                Guestbook.owner_id == current_user.id,
                Guestbook.is_deleted == False,
            )
        )
        if not result.scalar_one_or_none():
            raise HTTPException(status_code=404, detail="礼簿不存在")

    cached = _preview_cache.get(body.file_name)
    if cached is None:
        raise HTTPException(status_code=400, detail="请先进行导入预览")

    valid_indices = set(body.valid_row_indices)
    imported = 0
    skipped = 0

    try:
        for idx, cached_row in enumerate(cached["rows"]):
            if idx not in valid_indices:
                skipped += 1
                continue
            if not cached_row.get("is_valid") or cached_row.get("is_duplicate"):
                skipped += 1
                continue

            row_data = dict(cached_row["_parsed_data"])

            if body.import_type == "daily_record":
                record = DailyRecord(
                    uuid=new_uuid(),
                    owner_id=current_user.id,
                    source="daily",
                    created_by=current_user.username,
                    **row_data,
                )
                db.add(record)
                await db.flush()
                write_audit_log(
                    db, current_user.id, "daily_records", record.id, "CREATE",
                    operator_name=current_user.username,
                    note=f"批量导入自 {body.file_name}",
                )
            else:
                record_time = row_data.pop("record_time", None)
                record = GbRecord(
                    uuid=new_uuid(),
                    owner_id=current_user.id,
                    guestbook_id=body.guestbook_id,
                    check_status="unchecked",
                    record_time=record_time or datetime.now(),
                    **row_data,
                )
                db.add(record)
                await db.flush()
                write_audit_log(
                    db, current_user.id, "gb_records", record.id, "CREATE",
                    guestbook_id=body.guestbook_id,
                    operator_name=current_user.username,
                    note=f"批量导入自 {body.file_name}",
                )

            imported += 1

        await db.commit()
    except Exception as e:
        await db.rollback()
        raise HTTPException(status_code=500, detail=f"导入失败，已回滚：{str(e)}")

    _preview_cache.pop(body.file_name, None)

    return ImportCommitResponse(
        imported_count=imported,
        skipped_count=skipped,
        overwritten_count=0,
        errors=[],
    )
