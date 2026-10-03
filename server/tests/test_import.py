"""导入功能测试：整本礼簿 Excel 模板下载、预览与提交"""
import io

from openpyxl import Workbook


def _build_guestbook_excel(info: dict, rows: list[dict]) -> bytes:
    """按导入模板结构生成内存 xlsx：「礼簿信息」+「收礼明细」"""
    wb = Workbook()
    ws_info = wb.active
    ws_info.title = "礼簿信息"
    info_headers = ["礼簿名称", "事由", "礼簿日期", "备注"]
    for col, h in enumerate(info_headers, 1):
        ws_info.cell(row=1, column=col, value=h)
        ws_info.cell(row=2, column=col, value=info.get(h, ""))

    ws = wb.create_sheet("收礼明细")
    gb_headers = ["宾客姓名", "电话", "金额", "支付方式", "收银人", "记账时间", "备注"]
    for col, h in enumerate(gb_headers, 1):
        ws.cell(row=1, column=col, value=h)
    for r_idx, row in enumerate(rows, 2):
        for col, h in enumerate(gb_headers, 1):
            ws.cell(row=r_idx, column=col, value=row.get(h, ""))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestGuestbookTemplate:
    async def test_download_guestbook_template(self, client, admin_headers):
        resp = client.get("/api/import/templates/guestbook", headers=admin_headers)
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        # 模板可被 openpyxl 打开且包含两个工作表
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.content))
        assert "礼簿信息" in wb.sheetnames
        assert "收礼明细" in wb.sheetnames


class TestGuestbookImport:
    def _upload(self, client, headers, content: bytes, filename="整本礼簿.xlsx"):
        return client.post(
            "/api/import/preview?import_type=guestbook",
            files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=headers,
        )

    async def test_preview_valid(self, client, admin_headers):
        content = _build_guestbook_excel(
            {"礼簿名称": "张三婚礼", "事由": "婚礼", "礼簿日期": "2026-10-01"},
            [
                {"宾客姓名": "李四", "金额": "600.00", "支付方式": "现金"},
                {"宾客姓名": "赵六", "金额": "500.5", "支付方式": "微信", "电话": "13900139000"},
                {"宾客姓名": "", "金额": "abc"},  # 无效行
            ],
        )
        resp = self._upload(client, admin_headers, content)
        assert resp.status_code == 200
        data = resp.json()
        assert data["guestbook_info"]["is_valid"] is True
        assert data["valid_rows"] == 2
        assert data["invalid_rows"] == 1
        assert data["rows"][0]["is_valid"] is True
        assert data["rows"][2]["is_valid"] is False

    async def test_preview_invalid_info(self, client, admin_headers):
        content = _build_guestbook_excel(
            {"礼簿名称": "", "事由": "不存在的事由", "礼簿日期": "2026-13-40"},
            [{"宾客姓名": "李四", "金额": "600"}],
        )
        resp = self._upload(client, admin_headers, content)
        assert resp.status_code == 200
        info = resp.json()["guestbook_info"]
        assert info["is_valid"] is False
        assert any("礼簿名称" in e for e in info["errors"])
        assert any("事由" in e for e in info["errors"])
        assert any("日期" in e for e in info["errors"])

    async def test_preview_reject_csv(self, client, admin_headers):
        resp = client.post(
            "/api/import/preview?import_type=guestbook",
            files={"file": ("data.csv", "宾客姓名,金额\n李四,600", "text/csv")},
            headers=admin_headers,
        )
        assert resp.status_code == 400

    async def test_commit_creates_guestbook(self, client, admin_headers, db_session):
        content = _build_guestbook_excel(
            {"礼簿名称": "历史婚礼补录", "事由": "婚礼", "礼簿日期": "2026-09-20", "备注": "旧账本迁移"},
            [
                {"宾客姓名": "李四", "金额": "600.00", "支付方式": "现金", "记账时间": "2026-09-20 12:30"},
                {"宾客姓名": "赵六", "金额": "500.00", "支付方式": "微信"},
            ],
        )
        preview = self._upload(client, admin_headers, content)
        assert preview.status_code == 200
        file_name = preview.json()["file_name"]

        commit = client.post("/api/import/commit", json={
            "import_type": "guestbook",
            "file_name": file_name,
            "valid_row_indices": [0, 1],
        }, headers=admin_headers)
        assert commit.status_code == 200
        data = commit.json()
        assert data["imported_count"] == 2
        assert data["skipped_count"] == 0
        assert data["guestbook_name"] == "历史婚礼补录"
        gb_id = data["guestbook_id"]
        assert gb_id > 0

        # 礼簿已创建且为收尾中，明细免对账
        gb_resp = client.get(f"/api/guestbooks/{gb_id}", headers=admin_headers)
        assert gb_resp.status_code == 200
        gb = gb_resp.json()
        assert gb["name"] == "历史婚礼补录"
        assert gb["event_type"] == "wedding"
        assert gb["status"] == "closing"

        summary = client.get(f"/api/guestbooks/{gb_id}/records?page=1&page_size=50", headers=admin_headers)
        assert summary.status_code == 200
        records = summary.json()["items"]
        assert len(records) == 2
        assert all(r["check_status"] == "checked" for r in records)
        assert sum(r["amount"] for r in records) == 110000

    async def test_commit_requires_valid_info(self, client, admin_headers):
        content = _build_guestbook_excel(
            {"礼簿名称": "", "事由": "", "礼簿日期": ""},
            [{"宾客姓名": "李四", "金额": "600"}],
        )
        preview = self._upload(client, admin_headers, content)
        assert preview.status_code == 200
        file_name = preview.json()["file_name"]

        commit = client.post("/api/import/commit", json={
            "import_type": "guestbook",
            "file_name": file_name,
            "valid_row_indices": [0],
        }, headers=admin_headers)
        assert commit.status_code == 400

    async def test_commit_without_preview_fails(self, client, admin_headers):
        resp = client.post("/api/import/commit", json={
            "import_type": "guestbook",
            "file_name": "不存在的文件.xlsx",
            "valid_row_indices": [0],
        }, headers=admin_headers)
        assert resp.status_code == 400

    async def test_user_isolation(self, client, admin_headers, user_headers):
        content = _build_guestbook_excel(
            {"礼簿名称": "管理员婚礼", "事由": "婚礼", "礼簿日期": "2026-09-20"},
            [{"宾客姓名": "李四", "金额": "600"}],
        )
        preview = self._upload(client, admin_headers, content)
        file_name = preview.json()["file_name"]
        # 普通用户拿同一文件名提交：缓存不可见（按会话独立预览）
        commit = client.post("/api/import/commit", json={
            "import_type": "guestbook",
            "file_name": file_name,
            "valid_row_indices": [0],
        }, headers=user_headers)
        assert commit.status_code == 400
