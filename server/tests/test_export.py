import csv
import io


def _create_record(client, headers, person_name="张三", amount=50000, record_type="in",
                   event_type="wedding", event_date="2026-09-20", is_private=False):
    payload = {
        "type": record_type,
        "person_name": person_name,
        "amount": amount,
        "event_type": event_type,
        "event_date": event_date,
    }
    if is_private:
        payload["is_private"] = True
    return client.post("/api/daily-records/", json=payload, headers=headers)


def _create_guestbook(client, headers, name="测试礼簿"):
    return client.post("/api/guestbooks/", json={
        "name": name,
        "event_type": "wedding",
        "event_date": "2026-09-20",
    }, headers=headers)


def _create_gb_record(client, headers, gb_id, guest_name="李四", amount=30000):
    return client.post(f"/api/guestbooks/{gb_id}/records", json={
        "guest_name": guest_name,
        "amount": amount,
        "payment_method": "cash",
    }, headers=headers)


class TestDailyRecordsExportCsv:
    async def test_export_empty_csv(self, client, admin_headers):
        resp = client.get("/api/export/daily-records?format=csv", headers=admin_headers)
        assert resp.status_code == 200
        assert "text/csv" in resp.headers["content-type"]
        content = resp.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        assert len(rows) == 1  # header only

    async def test_export_csv_with_data(self, client, admin_headers, db_session):
        _create_record(client, admin_headers, person_name="张三", amount=50000)
        _create_record(client, admin_headers, person_name="李四", amount=30000)

        resp = client.get("/api/export/daily-records?format=csv", headers=admin_headers)
        assert resp.status_code == 200
        content = resp.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        assert len(rows) == 3  # header + 2 records

    async def test_export_csv_filter_by_type(self, client, admin_headers, db_session):
        _create_record(client, admin_headers, person_name="张三", record_type="in")
        _create_record(client, admin_headers, person_name="李四", record_type="out")

        resp = client.get("/api/export/daily-records?format=csv&type=in", headers=admin_headers)
        content = resp.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        assert len(rows) == 2  # header + 1 record

    async def test_export_csv_filter_by_date(self, client, admin_headers, db_session):
        _create_record(client, admin_headers, person_name="张三", event_date="2026-09-01")
        _create_record(client, admin_headers, person_name="李四", event_date="2026-09-20")

        resp = client.get(
            "/api/export/daily-records?format=csv&start_date=2026-09-15&end_date=2026-09-25",
            headers=admin_headers,
        )
        content = resp.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        assert len(rows) == 2  # header + 1 record

    async def test_export_csv_exclude_private(self, client, admin_headers, db_session):
        _create_record(client, admin_headers, person_name="张三")
        _create_record(client, admin_headers, person_name="李四", is_private=True)

        resp = client.get(
            "/api/export/daily-records?format=csv&include_private=false",
            headers=admin_headers,
        )
        content = resp.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        assert len(rows) == 2  # header + 1 non-private record


class TestDailyRecordsExportExcel:
    async def test_export_empty_xlsx(self, client, admin_headers):
        resp = client.get("/api/export/daily-records?format=xlsx", headers=admin_headers)
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        assert len(resp.content) > 0

    async def test_export_xlsx_with_data(self, client, admin_headers, db_session):
        _create_record(client, admin_headers, person_name="张三")

        resp = client.get("/api/export/daily-records?format=xlsx", headers=admin_headers)
        assert resp.status_code == 200

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        assert ws.title == "日常记录"
        assert ws.cell(row=2, column=1).value == "张三"

    async def test_export_default_format_is_xlsx(self, client, admin_headers):
        resp = client.get("/api/export/daily-records", headers=admin_headers)
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]

    async def test_export_invalid_format_rejected(self, client, admin_headers):
        resp = client.get("/api/export/daily-records?format=pdf", headers=admin_headers)
        assert resp.status_code == 422


class TestGuestbookExport:
    async def test_export_guestbook_empty(self, client, admin_headers, db_session):
        gb_id = _create_guestbook(client, admin_headers).json()["id"]

        resp = client.get(f"/api/export/guestbooks/{gb_id}", headers=admin_headers)
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.content))
        sheet_names = wb.sheetnames
        assert "收礼记录" in sheet_names
        assert "汇总" in sheet_names

    async def test_export_guestbook_with_records(self, client, admin_headers, db_session):
        gb_id = _create_guestbook(client, admin_headers).json()["id"]
        _create_gb_record(client, admin_headers, gb_id, guest_name="宾客A", amount=50000)
        _create_gb_record(client, admin_headers, gb_id, guest_name="宾客B", amount=30000)

        resp = client.get(f"/api/export/guestbooks/{gb_id}", headers=admin_headers)
        assert resp.status_code == 200

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb["收礼记录"]
        assert ws.max_row == 3  # header + 2 records

        ws_summary = wb["汇总"]
        assert ws_summary.cell(row=1, column=2).value == "测试礼簿"

    async def test_export_guestbook_not_found(self, client, admin_headers):
        resp = client.get("/api/export/guestbooks/9999", headers=admin_headers)
        assert resp.status_code == 404

    async def test_export_guestbook_data_isolation(self, client, admin_headers, another_user_headers, db_session):
        gb_id = _create_guestbook(client, admin_headers).json()["id"]

        resp = client.get(f"/api/export/guestbooks/{gb_id}", headers=another_user_headers)
        assert resp.status_code == 404


class TestExportDataIsolation:
    async def test_user_cannot_see_other_records(self, client, admin_headers, another_user_headers, db_session):
        _create_record(client, admin_headers, person_name="张三")

        resp = client.get("/api/export/daily-records?format=csv", headers=another_user_headers)
        content = resp.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        assert len(rows) == 1  # header only
